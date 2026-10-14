#!/usr/bin/env python3
"""
rightsizing_check.py

Validate AWS rightsizing recommendations (the Cost Optimization Hub /
Compute Optimizer export your FinOps team sent) against real utilization,
week by week, across every account in the sheet.

For each row (Ec2AutoScalingGroup or Ec2Instance) it pulls from CloudWatch:
  - CPU %                 (AWS/EC2 CPUUtilization)
  - Memory %              (CWAgent "Memory % Committed Bytes In Use" on Windows,
                           "mem_used_percent" on Linux; needs the CloudWatch agent)
  - Network in/out Mbps   (AWS/EC2 NetworkIn / NetworkOut)
  - EBS MB/s and IOPS     (AWS/EC2 EBSRead/WriteBytes, EBSRead/WriteOps; Nitro)

It rolls them up per week (Mon-Sun, UTC): peak, p95, average. Then it scales
the peaks onto the recommended instance type (vCPU and RAM ratio) and gives a
verdict per resource.

Output: one Excel workbook
  Summary : one row per resource; current vs recommended size, overall peaks,
            projected utilization on the recommended type, verdict + reason
  Weekly  : one row per resource per week; actual peak/p95/avg per metric
  Notes   : how to read the numbers

Usage
  pip install boto3 pandas openpyxl

  # One AWS CLI profile per account
  python rightsizing_check.py usmg-cost-saving-opportunities.xlsx \
      --profile-map 123456789=usmg-dev 111122223333=usmg-qa \
      --weeks 4

  # Or assume the same read-only role in every account from one base profile
  python rightsizing_check.py usmg-cost-saving-opportunities.xlsx \
      --base-profile my-sso-profile --role-name ReadOnlyRole --weeks 6

Required permissions (read-only):
  cloudwatch:GetMetricData, ec2:DescribeInstances, ec2:DescribeInstanceTypes,
  autoscaling:DescribeAutoScalingGroups, sts:GetCallerIdentity
  (+ sts:AssumeRole when --role-name is used)
"""
import argparse
import datetime as dt
import re
import sys
from collections import defaultdict

import boto3
import pandas as pd
from botocore.exceptions import ClientError

# ---------------- verdict thresholds (tune to your team's comfort) ----------
CPU_P95_LIMIT = 70    # projected p95 CPU % on the recommended type -> KEEP if above
CPU_PEAK_LIMIT = 90   # projected peak CPU % on the recommended type -> BORDERLINE if above
MEM_PEAK_LIMIT = 85   # projected peak memory % -> KEEP if above
MEM_WARN = 75         # projected peak memory % -> BORDERLINE if above
NET_LIMIT = 0.8       # fraction of the target's baseline network bandwidth

WIN_MEM = "Memory % Committed Bytes In Use"
LINUX_MEM = "mem_used_percent"

# key, CloudWatch metric, statistic
EC2_METRICS = [
    ("cpumax", "CPUUtilization", "Maximum"),
    ("cpuavg", "CPUUtilization", "Average"),
    ("netin", "NetworkIn", "Maximum"),
    ("netout", "NetworkOut", "Maximum"),
    ("ebsrb", "EBSReadBytes", "Maximum"),
    ("ebswb", "EBSWriteBytes", "Maximum"),
    ("ebsro", "EBSReadOps", "Maximum"),
    ("ebswo", "EBSWriteOps", "Maximum"),
]


# ---------------------------------------------------------------- helpers ---
def log(msg):
    print(msg, file=sys.stderr, flush=True)


def pick(row, *names, default=None):
    """Case-insensitive column lookup with fallbacks."""
    lower = {str(k).strip().lower(): k for k in row.index}
    for n in names:
        k = lower.get(n.lower())
        if k is not None and pd.notna(row[k]) and str(row[k]).strip() != "":
            return row[k]
    return default


def parse_type(text):
    """'r6i.xlarge', 'm7i-flex.large', '... r8i-flex.large ...' -> type string."""
    if text is None:
        return None
    m = re.search(r"\b([a-z][a-z0-9]*(?:-flex)?\.[0-9a-z]+)\b", str(text).lower())
    return m.group(1) if m else None


def is_asg(rtype):
    return "autoscaling" in str(rtype).lower()


class Sessions:
    """Returns a boto3 Session for an account (profile map, assume-role, or default)."""

    def __init__(self, args):
        self.args = args
        self.cache = {}
        self.profile_map = {}
        for p in args.profile_map or []:
            acct, prof = p.split("=", 1)
            self.profile_map[acct.strip().zfill(12)] = prof.strip()

    def _base(self):
        a = self.args
        return boto3.Session(profile_name=a.base_profile) if a.base_profile else boto3.Session()

    def get(self, account):
        if account in self.cache:
            return self.cache[account]
        if account in self.profile_map:
            s = boto3.Session(profile_name=self.profile_map[account])
        elif self.args.role_name:
            c = self._base().client("sts").assume_role(
                RoleArn=f"arn:aws:iam::{account}:role/{self.args.role_name}",
                RoleSessionName="rightsizing-check",
            )["Credentials"]
            s = boto3.Session(
                aws_access_key_id=c["AccessKeyId"],
                aws_secret_access_key=c["SecretAccessKey"],
                aws_session_token=c["SessionToken"],
            )
        else:
            s = self._base()
        actual = s.client("sts").get_caller_identity()["Account"]
        if actual != account:
            raise RuntimeError(
                f"credentials point to account {actual}, not {account} "
                f"(add --profile-map {account}=<profile> or use --role-name)"
            )
        self.cache[account] = s
        return s


_spec_cache = {}


def type_specs(ec2, itype):
    if not itype:
        return None
    if itype in _spec_cache:
        return _spec_cache[itype]
    try:
        t = ec2.describe_instance_types(InstanceTypes=[itype])["InstanceTypes"][0]
    except (ClientError, IndexError):
        _spec_cache[itype] = None
        return None
    cards = t.get("NetworkInfo", {}).get("NetworkCards", [])
    ebs = t.get("EbsInfo", {}).get("EbsOptimizedInfo", {})
    spec = {
        "vcpu": t["VCpuInfo"]["DefaultVCpus"],
        "mem_gib": round(t["MemoryInfo"]["SizeInMiB"] / 1024, 1),
        "net_base_gbps": sum(c.get("BaselineBandwidthInGbps", 0) for c in cards) or None,
        "ebs_base_mbs": ebs.get("BaselineThroughputInMBps"),
        "ebs_base_iops": ebs.get("BaselineIops"),
    }
    _spec_cache[itype] = spec
    return spec


def current_instance_ids(session, region, rtype, rid):
    if is_asg(rtype):
        asg = session.client("autoscaling", region_name=region)
        groups = asg.describe_auto_scaling_groups(AutoScalingGroupNames=[rid])["AutoScalingGroups"]
        return [i["InstanceId"] for i in groups[0]["Instances"]] if groups else []
    return [rid]


def sample_interval(ec2, ids):
    """60s if detailed monitoring is on, else 300s. Needed to turn byte/ops counts into rates."""
    if not ids:
        return 300
    try:
        res = ec2.describe_instances(InstanceIds=ids[:100])
        for r in res["Reservations"]:
            for i in r["Instances"]:
                if i.get("Monitoring", {}).get("State") == "enabled":
                    return 60
    except ClientError:
        pass
    return 300


# ------------------------------------------------------------- CloudWatch ---
def metric_q(qid, name, dims, stat, period):
    return {
        "Id": qid,
        "MetricStat": {
            "Metric": {"Namespace": "AWS/EC2", "MetricName": name, "Dimensions": dims},
            "Period": period,
            "Stat": stat,
        },
        "ReturnData": True,
    }


def ec2_queries(dim_sets, period):
    q = []
    for k, d in enumerate(dim_sets):
        for key, name, stat in EC2_METRICS:
            q.append(metric_q(f"{key}_{k}", name, [d], stat, period))
    return q


def mem_queries(rtype, rid, ids, period):
    """Search CWAgent memory metrics by ASG name (if the agent appends it) and by instance id."""
    filters = [f'AutoScalingGroupName="{rid}"'] if is_asg(rtype) else []
    filters += [f'InstanceId="{i}"' for i in ids[:40]]
    q, n = [], 0
    for f in filters:
        for m in (WIN_MEM, LINUX_MEM):
            expr = f"SEARCH('Namespace=\"CWAgent\" MetricName=\"{m}\" {f}', 'Maximum', {period})"
            q.append({"Id": f"mem_{n}", "Expression": expr, "ReturnData": True})
            n += 1
    return q


def fetch(cw, queries, start, end):
    """Run GetMetricData, return {metric_key: combined Series (max across series)}."""
    raw = defaultdict(list)
    for i in range(0, len(queries), 500):
        pages = cw.get_paginator("get_metric_data").paginate(
            MetricDataQueries=queries[i:i + 500],
            StartTime=start, EndTime=end, ScanBy="TimestampAscending",
        )
        for page in pages:
            for r in page["MetricDataResults"]:
                if r["Timestamps"]:
                    key = r["Id"].split("_")[0]
                    raw[key].append(pd.Series(r["Values"], index=pd.to_datetime(r["Timestamps"], utc=True)))
    out = {}
    for key, parts in raw.items():
        s = pd.concat(parts)
        out[key] = s.groupby(level=0).max().sort_index()
    return out


# --------------------------------------------------------------- analysis ---
def build_frame(series, interval):
    """Turn raw metric series into one 5-min frame with human units."""
    empty = pd.Series(dtype=float, index=pd.DatetimeIndex([], tz="UTC"))

    def get(k):
        return series.get(k, empty)

    df = pd.DataFrame({
        "cpu_max": get("cpumax"),
        "cpu_avg": get("cpuavg"),
        "mem": get("mem"),
        "net_in_mbps": get("netin") * 8 / interval / 1e6,
        "net_out_mbps": get("netout") * 8 / interval / 1e6,
        "ebs_mbs": get("ebsrb").add(get("ebswb"), fill_value=0) / interval / 1e6,
        "ebs_iops": get("ebsro").add(get("ebswo"), fill_value=0) / interval,
    })
    df.index = pd.to_datetime(df.index, utc=True)
    return df.sort_index()


def q95(x):
    x = x.dropna()
    return x.quantile(0.95) if len(x) else float("nan")


def weekly_rollup(df, period):
    if df.empty:
        return pd.DataFrame()
    wk = df.index.tz_convert(None).to_period("W").start_time.date
    g = df.groupby(wk)
    out = pd.DataFrame({
        "cpu_peak_%": g["cpu_max"].max(),
        "cpu_p95_%": g["cpu_avg"].agg(q95),
        "cpu_avg_%": g["cpu_avg"].mean(),
        "mem_peak_%": g["mem"].max(),
        "mem_p95_%": g["mem"].agg(q95),
        "mem_avg_%": g["mem"].mean(),
        "net_in_peak_mbps": g["net_in_mbps"].max(),
        "net_out_peak_mbps": g["net_out_mbps"].max(),
        "ebs_peak_MBps": g["ebs_mbs"].max(),
        "ebs_peak_iops": g["ebs_iops"].max(),
        "hours_of_data": g["cpu_avg"].count() * period / 3600,
    })
    out.index.name = "week_starting"
    return out.reset_index()


def verdict(r):
    if pd.isna(r.get("cpu_peak_%")):
        return "NO DATA", "No CloudWatch CPU data in the window (resource idle/terminated?)"
    cur_mem, rec_mem = r.get("cur_mem_gib"), r.get("rec_mem_gib")
    if rec_mem is None or r.get("rec_vcpu") is None:
        return "CHECK MANUALLY", "Could not look up the recommended instance type"

    keep, warn = [], []
    if r["proj_cpu_p95_%"] > CPU_P95_LIMIT:
        keep.append(f"p95 CPU would be ~{r['proj_cpu_p95_%']:.0f}%")
    elif r["proj_cpu_peak_%"] > CPU_PEAK_LIMIT:
        warn.append(f"peak CPU would be ~{r['proj_cpu_peak_%']:.0f}%")

    if pd.isna(r.get("mem_peak_%")):
        if rec_mem < cur_mem:
            return ("NEED MEMORY DATA",
                    f"Target has {rec_mem} GiB vs {cur_mem} GiB today and no CloudWatch agent "
                    f"memory metric was found; the recommendation is CPU-only")
    else:
        if r["proj_mem_peak_%"] > MEM_PEAK_LIMIT:
            keep.append(f"peak memory would be ~{r['proj_mem_peak_%']:.0f}%")
        elif r["proj_mem_peak_%"] > MEM_WARN:
            warn.append(f"peak memory would be ~{r['proj_mem_peak_%']:.0f}%")

    base = r.get("rec_net_base_gbps")
    net_peak = max(r.get("net_in_peak_mbps") or 0, r.get("net_out_peak_mbps") or 0)
    if base and net_peak > NET_LIMIT * base * 1000:
        keep.append(f"network peak {net_peak:.0f} Mbps vs {base} Gbps baseline")

    if keep:
        return "KEEP CURRENT", "; ".join(keep + warn)
    if warn:
        return "BORDERLINE", "; ".join(warn)
    return "SAFE TO DOWNSIZE", "Peaks fit the recommended type with headroom"


def analyze_row(row, sessions, start, end, period):
    account = str(pick(row, "accountId")).split(".")[0].zfill(12)
    region = pick(row, "region")
    rtype = pick(row, "currentResourceType", default="Ec2Instance")
    rid = str(pick(row, "resourceId")).split("/")[-1]
    cur_t = parse_type(pick(row, "currentResourceSummary", "currentResourceDetails"))
    rec_t = parse_type(pick(row, "recommendedResourceSummary", "recommendedResourceDetails"))

    base = {
        "accountId": account,
        "accountName": pick(row, "accountName", default=""),
        "region": region,
        "resourceType": rtype,
        "resourceId": rid,
        "current_type": cur_t,
        "recommended_type": rec_t,
        "est_monthly_savings_$": pick(row, "estimatedMonthlySavings", "estimatedMonthlySaving"),
    }

    s = sessions.get(account)
    ec2 = s.client("ec2", region_name=region)
    cw = s.client("cloudwatch", region_name=region)

    ids = current_instance_ids(s, region, rtype, rid)
    interval = sample_interval(ec2, ids)

    dim_key = "AutoScalingGroupName" if is_asg(rtype) else "InstanceId"
    series = fetch(cw, ec2_queries([{"Name": dim_key, "Value": rid}], period)
                   + mem_queries(rtype, rid, ids, period), start, end)
    if "cpumax" not in series and is_asg(rtype) and ids:
        # ASG-level aggregate not available -> fall back to the current instances
        fallback = fetch(cw, ec2_queries([{"Name": "InstanceId", "Value": i} for i in ids[:40]], period),
                         start, end)
        series.update(fallback)

    df = build_frame(series, interval)
    weekly = weekly_rollup(df, period)
    if not weekly.empty:
        for k, v in reversed(list(base.items())[:5]):
            weekly.insert(0, k, v)

    cur = type_specs(ec2, cur_t) or {}
    rec = type_specs(ec2, rec_t) or {}

    summ = dict(base)
    summ.update({
        "cur_vcpu": cur.get("vcpu"), "cur_mem_gib": cur.get("mem_gib"),
        "rec_vcpu": rec.get("vcpu"), "rec_mem_gib": rec.get("mem_gib"),
        "rec_net_base_gbps": rec.get("net_base_gbps"),
        "weeks_of_data": int((weekly["hours_of_data"] > 0).sum()) if not weekly.empty else 0,
        "cpu_peak_%": df["cpu_max"].max() if not df.empty else float("nan"),
        "cpu_p95_%": q95(df["cpu_avg"]) if not df.empty else float("nan"),
        "mem_peak_%": df["mem"].max() if not df.empty else float("nan"),
        "mem_p95_%": q95(df["mem"]) if not df.empty else float("nan"),
        "net_in_peak_mbps": df["net_in_mbps"].max() if not df.empty else float("nan"),
        "net_out_peak_mbps": df["net_out_mbps"].max() if not df.empty else float("nan"),
        "ebs_peak_MBps": df["ebs_mbs"].max() if not df.empty else float("nan"),
        "ebs_peak_iops": df["ebs_iops"].max() if not df.empty else float("nan"),
        "weeks_cpu_peak_over_80%": int((weekly["cpu_peak_%"] > 80).sum()) if not weekly.empty else 0,
        "weeks_mem_peak_over_80%": int((weekly["mem_peak_%"] > 80).sum()) if not weekly.empty else 0,
    })

    cpu_ratio = (cur["vcpu"] / rec["vcpu"]) if cur.get("vcpu") and rec.get("vcpu") else float("nan")
    mem_ratio = (cur["mem_gib"] / rec["mem_gib"]) if cur.get("mem_gib") and rec.get("mem_gib") else float("nan")
    summ["proj_cpu_peak_%"] = summ["cpu_peak_%"] * cpu_ratio
    summ["proj_cpu_p95_%"] = summ["cpu_p95_%"] * cpu_ratio
    summ["proj_mem_peak_%"] = summ["mem_peak_%"] * mem_ratio
    summ["proj_mem_p95_%"] = summ["mem_p95_%"] * mem_ratio
    summ["verdict"], summ["reason"] = verdict(summ)
    return summ, weekly


# ------------------------------------------------------------------ output --
NOTES = [
    ("What this is", "Real CloudWatch utilization, per week, for every resource the FinOps "
                     "export recommends downsizing, plus how the busiest moments would look on "
                     "the recommended instance type."),
    ("Peak", "Highest 5-minute value in the window (busiest instance in the ASG)."),
    ("p95", "95% of 5-minute averages were at or below this. Good 'normal busy' indicator."),
    ("proj_* columns", "Current utilization scaled to the recommended type: CPU x (current vCPU / "
                       "target vCPU), memory x (current GiB / target GiB). Over 100% means the "
                       "target would saturate. Newer generations (m7i, r8i) are ~10-20% faster per "
                       "vCPU, so CPU projections are slightly pessimistic."),
    ("Memory", "Comes only from the CloudWatch agent. On Windows it is 'Memory % Committed Bytes "
               "In Use' (includes page file), so it is a conservative proxy for RAM. If it is "
               "blank, AWS's recommendation did not consider memory either."),
    ("Network / EBS", "Busiest 5-minute (or 1-minute with detailed monitoring) rate. Short bursts "
                      "inside that interval are averaged out."),
    ("Verdict", f"KEEP CURRENT: projected p95 CPU > {CPU_P95_LIMIT}% or peak memory > "
                f"{MEM_PEAK_LIMIT}% or network > {int(NET_LIMIT*100)}% of target baseline. "
                f"BORDERLINE: projected peak CPU > {CPU_PEAK_LIMIT}% or memory > {MEM_WARN}%. "
                f"NEED MEMORY DATA: target has less RAM and no memory metric exists."),
    ("Flex types", "m7i-flex / r8i-flex deliver full CPU most of the time but are not meant for "
                   "sustained 100% CPU; treat BORDERLINE on a flex target as KEEP."),
]

FILLS = {
    "SAFE TO DOWNSIZE": "C6EFCE",
    "BORDERLINE": "FFEB9C",
    "KEEP CURRENT": "FFC7CE",
    "NEED MEMORY DATA": "FCE4D6",
    "CHECK MANUALLY": "D9D9D9",
    "NO DATA": "D9D9D9",
    "ERROR": "D9D9D9",
}


def write_excel(path, summary, weekly):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        weekly.to_excel(xw, sheet_name="Weekly", index=False)
        pd.DataFrame(NOTES, columns=["Topic", "Explanation"]).to_excel(xw, sheet_name="Notes", index=False)

        for name in ("Summary", "Weekly", "Notes"):
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10),
                                                                       90 if name == "Notes" else 45)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0"

        ws = xw.sheets["Summary"]
        headers = [c.value for c in ws[1]]
        if "verdict" in headers:
            vcol = headers.index("verdict") + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=vcol)
                color = FILLS.get(cell.value)
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)
                    cell.font = Font(bold=True)


# -------------------------------------------------------------------- main --
def main():
    ap = argparse.ArgumentParser(description="Validate EC2 rightsizing recommendations with weekly CloudWatch peaks.")
    ap.add_argument("input", help="FinOps export (.xlsx or .csv)")
    ap.add_argument("--sheet", default=0, help="sheet name/index for .xlsx (default: first)")
    ap.add_argument("--weeks", type=int, default=4, help="full weeks to look back (default 4)")
    ap.add_argument("--profile-map", nargs="*", metavar="ACCOUNT=PROFILE",
                    help="AWS CLI profile per account, e.g. 123456789=usmg-dev")
    ap.add_argument("--base-profile", help="profile to use directly or to assume roles from")
    ap.add_argument("--role-name", help="role to assume in each account (e.g. ReadOnlyRole)")
    ap.add_argument("--accounts", nargs="*", help="only these account IDs")
    ap.add_argument("--output", default=None, help="output .xlsx (default: rightsizing_check_<date>.xlsx)")
    args = ap.parse_args()

    if args.input.lower().endswith(".csv"):
        src = pd.read_csv(args.input, dtype=str)
    else:
        sheet = int(args.sheet) if str(args.sheet).isdigit() else args.sheet
        src = pd.read_excel(args.input, sheet_name=sheet, dtype=str)

    rows = []
    seen = set()
    for _, row in src.iterrows():
        rtype = str(pick(row, "currentResourceType", default=""))
        action = str(pick(row, "actionType", default="Rightsize"))
        if rtype not in ("Ec2AutoScalingGroup", "Ec2Instance") or action.lower() != "rightsize":
            continue
        acct = str(pick(row, "accountId", default="")).split(".")[0].zfill(12)
        if args.accounts and acct not in {a.zfill(12) for a in args.accounts}:
            continue
        key = (acct, pick(row, "region"), pick(row, "resourceId"))
        if key in seen:
            continue
        seen.add(key)
        rows.append(row)

    if not rows:
        sys.exit("No Ec2AutoScalingGroup / Ec2Instance 'Rightsize' rows found in the input.")

    now = dt.datetime.now(dt.timezone.utc)
    monday = (now - dt.timedelta(days=now.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    start = monday - dt.timedelta(weeks=args.weeks)
    period = 300 if args.weeks <= 8 else 3600   # 5-minute data is kept for 63 days
    log(f"{len(rows)} resources, window {start:%Y-%m-%d} -> {now:%Y-%m-%d %H:%M} UTC, period {period}s")

    sessions = Sessions(args)
    summaries, weeklies = [], []
    for n, row in enumerate(rows, 1):
        rid = pick(row, "resourceId")
        log(f"[{n}/{len(rows)}] {pick(row, 'accountId')} {pick(row, 'region')} {rid}")
        try:
            summ, weekly = analyze_row(row, sessions, start, now, period)
            summaries.append(summ)
            if not weekly.empty:
                weeklies.append(weekly)
        except Exception as e:  # keep going; record the failure in the report
            log(f"   ! {e}")
            summaries.append({
                "accountId": str(pick(row, "accountId")), "region": pick(row, "region"),
                "resourceId": rid, "verdict": "ERROR", "reason": str(e)[:300],
            })

    summary = pd.DataFrame(summaries)
    order = ["accountId", "accountName", "region", "resourceType", "resourceId",
             "current_type", "recommended_type", "verdict", "reason", "est_monthly_savings_$",
             "cur_vcpu", "cur_mem_gib", "rec_vcpu", "rec_mem_gib", "weeks_of_data",
             "cpu_peak_%", "cpu_p95_%", "proj_cpu_peak_%", "proj_cpu_p95_%",
             "mem_peak_%", "mem_p95_%", "proj_mem_peak_%", "proj_mem_p95_%",
             "weeks_cpu_peak_over_80%", "weeks_mem_peak_over_80%",
             "net_in_peak_mbps", "net_out_peak_mbps", "rec_net_base_gbps",
             "ebs_peak_MBps", "ebs_peak_iops"]
    summary = summary.reindex(columns=[c for c in order if c in summary.columns]
                              + [c for c in summary.columns if c not in order])
    weekly = pd.concat(weeklies, ignore_index=True) if weeklies else pd.DataFrame()

    out = args.output or f"rightsizing_check_{now:%Y-%m-%d}.xlsx"
    write_excel(out, summary, weekly)
    log(f"\nWrote {out}")
    if "verdict" in summary:
        log(summary["verdict"].value_counts().to_string())


if __name__ == "__main__":
    main()
